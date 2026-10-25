# -*- coding: utf-8 -*-
"""
Génère la documentation professionnelle du projet d'archivage DGRAD.
Sortie : docs/Documentation_Projet_Archivage_DGRAD.docx
"""

import math
import os
from datetime import date

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCS_DIR = os.path.join(BASE_DIR, "docs")
ASSETS_DIR = os.path.join(DOCS_DIR, "assets")
OUTPUT = os.path.join(DOCS_DIR, "Documentation_Projet_Archivage_DGRAD.docx")
LOGO = os.path.join(BASE_DIR, "Frontend", "src", "Images", "Logo.png")

os.makedirs(ASSETS_DIR, exist_ok=True)

NAVY = RGBColor(0x09, 0x18, 0x27)
NAVY_HEX = "091827"
GOLD = RGBColor(0xC7, 0x9A, 0x27)
GOLD_HEX = "C79A27"
SLATE = RGBColor(0x47, 0x55, 0x69)
SLATE_LIGHT = RGBColor(0x64, 0x74, 0x8B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_BG = "F4F6FA"
LIGHT_GOLD = "FBF3D9"
BORDER = "D9DEE7"
GREEN = RGBColor(0x05, 0x96, 0x69)
RED = RGBColor(0xDC, 0x26, 0x26)


# ---------------------------------------------------------------------------
# Diagrammes
# ---------------------------------------------------------------------------

def _font(size, bold=False):
    candidates = [
        "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf",
        "C:/Windows/Fonts/segoeui.ttf",
    ]
    for path in candidates:
        if os.path.exists(path):
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()


def _rounded_box(draw, box, title, lines, fill, border, title_color="#091827"):
    x, y, w, h = box
    draw.rounded_rectangle([x, y, x + w, y + h], radius=18, fill=fill, outline=border, width=3)
    draw.text((x + 24, y + 18), title, font=_font(25, True), fill=title_color)
    for index, line in enumerate(lines):
        draw.text((x + 24, y + 58 + index * 27), line, font=_font(18), fill="#475569")


def _arrow(draw, start, end, color, label=None):
    draw.line([start, end], fill=color, width=4)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    length = 18
    for delta in (-0.45, 0.45):
        x = end[0] - length * math.cos(angle + delta)
        y = end[1] - length * math.sin(angle + delta)
        draw.line([end, (x, y)], fill=color, width=4)
    if label:
        mx = (start[0] + end[0]) / 2
        my = (start[1] + end[1]) / 2
        draw.rounded_rectangle([mx - 65, my - 18, mx + 65, my + 18], radius=8, fill="#FFFFFF")
        draw.text((mx - 58, my - 12), label, font=_font(15, True), fill="#64748B")


def build_architecture_image(path):
    width, height = 1800, 1000
    image = Image.new("RGB", (width, height), "#F7F8FA")
    draw = ImageDraw.Draw(image)

    draw.text((60, 42), "Architecture applicative", font=_font(38, True), fill="#091827")
    draw.text((60, 92), "DGRAD — Plateforme de gestion des archives et des titres de perception", font=_font(20), fill="#64748B")

    boxes = {
        "users": ((70, 210, 320, 130), "Utilisateurs", ["Navigateur web", "Agents DGRAD"], "#FBF3D9", "#D9B145"),
        "front": ((470, 210, 360, 130), "Frontend React", ["Vite + Tailwind", "Permissions & périmètre"], "#EAF2FF", "#93B4E8"),
        "api": ((910, 210, 360, 130), "API Laravel 8", ["Sanctum + RBAC", "Contrôleurs métier"], "#091827", "#091827"),
        "db": ((1370, 210, 360, 130), "MySQL 8", ["bdarchive_titres_perception", "Données & audit"], "#E8F7F1", "#7CC7AA"),
        "scan": ((470, 600, 360, 130), "Scanner Windows", ["WindowScan", "Numérisation PDF"], "#EEF2F7", "#AAB7C7"),
        "ocr": ((910, 600, 360, 130), "Service OCR Flask", ["Tesseract / pdfplumber", "Extraction de texte"], "#FFF4DC", "#E2BD63"),
        "files": ((1370, 600, 360, 130), "Stockage fichiers", ["storage/app/avatar", "scans & PDF"], "#F2EDFB", "#BBA6E8"),
    }

    for _, (box, title, lines, fill, border) in boxes.items():
        title_color = "#FFFFFF" if fill == "#091827" else "#091827"
        _rounded_box(draw, box, title, lines, fill, border, title_color)

    _arrow(draw, (390, 275), (470, 275), "#C79A27", "HTTPS")
    _arrow(draw, (830, 275), (910, 275), "#C79A27", "JSON")
    _arrow(draw, (1270, 275), (1370, 275), "#64748B", "SQL")
    _arrow(draw, (1090, 600), (1090, 340), "#C79A27", "OCR")
    _arrow(draw, (830, 665), (910, 665), "#64748B", "PDF")
    _arrow(draw, (1270, 665), (1370, 665), "#64748B", "Stockage")
    _arrow(draw, (1180, 340), (1180, 600), "#C79A27", "Texte")

    draw.line([(60, 900), (1740, 900)], fill="#E2E8F0", width=2)
    draw.text((60, 930), "Flux principal : navigateur → API Laravel → MySQL / OCR / stockage", font=_font(17), fill="#64748B")
    draw.text((1390, 930), "Version 1.0 — 2026", font=_font(17, True), fill="#C79A27")
    image.save(path, "PNG")


def build_data_model_image(path):
    width, height = 1800, 1050
    image = Image.new("RGB", (width, height), "#F7F8FA")
    draw = ImageDraw.Draw(image)

    draw.text((60, 42), "Cartographie des domaines de données", font=_font(38, True), fill="#091827")
    draw.text((60, 92), "Vue simplifiée des relations entre les principaux groupes de tables", font=_font(20), fill="#64748B")

    boxes = {
        "org": ((70, 190, 430, 175), "Organisation", ["directions", "divisions", "centre_ordonnancements"], "#EAF2FF", "#93B4E8"),
        "rbac": ((70, 450, 430, 190), "Utilisateurs & RBAC", ["users · roles · permissions", "user_role · role_permission", "division_user"], "#FBF3D9", "#D9B145"),
        "ref": ((620, 190, 500, 190), "Référentiels", ["types_documents · classeurs", "emplacements · mots_cles", "assujettis · articles_budgetaires", "categories_titres"], "#E8F7F1", "#7CC7AA"),
        "arch": ((1270, 190, 460, 190), "Archivage ordinaire", ["archives", "archive_mot_cle", "fichiers & OCR"], "#EAF2FF", "#93B4E8"),
        "titre": ((1270, 470, 460, 220), "Titres de perception", ["titres_perception", "titre_mot_cle", "fichiers & OCR"], "#F2EDFB", "#BBA6E8"),
        "ocr": ((620, 470, 500, 190), "Fichiers & OCR", ["fichiers_documents", "traitements_ocr", "texte_extrait"], "#FFF4DC", "#E2BD63"),
        "audit": ((620, 760, 500, 160), "Audit & technique", ["audit_logs", "jobs · failed_jobs", "migrations · password_resets"], "#EEF2F7", "#AAB7C7"),
    }

    for _, (box, title, lines, fill, border) in boxes.items():
        _rounded_box(draw, box, title, lines, fill, border)

    _arrow(draw, (500, 275), (620, 275), "#64748B", "1-N")
    _arrow(draw, (500, 520), (620, 520), "#64748B", "accès")
    _arrow(draw, (1120, 275), (1270, 275), "#64748B", "classement")
    _arrow(draw, (1120, 560), (1270, 560), "#64748B", "classement")
    _arrow(draw, (1500, 380), (1500, 470), "#64748B", "fichiers")
    _arrow(draw, (1120, 760), (620, 845), "#64748B", "trace")
    _arrow(draw, (500, 540), (620, 800), "#64748B", "journal")

    draw.line([(60, 960), (1740, 960)], fill="#E2E8F0", width=2)
    draw.text((60, 985), "Les entités archives et titres partagent les mêmes référentiels, fichiers et mécanismes d’OCR.", font=_font(17), fill="#64748B")
    image.save(path, "PNG")


# ---------------------------------------------------------------------------
# Helpers Word
# ---------------------------------------------------------------------------

def set_cell_shading(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def set_table_borders(table, color=BORDER, size="4"):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)
        borders.append(element)
    tbl_pr.append(borders)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_field(paragraph, instruction, placeholder=""):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    run._r.append(text)
    run._r.append(end)


def add_toc(doc):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(10)
    add_field(paragraph, 'TOC \\o "1-3" \\h \\z \\u', "Mettre à jour le sommaire dans Word (F9).")


def add_page_number_footer(section):
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.text = ""
    paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(16.2), WD_TAB_ALIGNMENT.RIGHT)

    left = paragraph.add_run("Documentation plateforme archives & titres — Version 1.0")
    left.font.size = Pt(8)
    left.font.color.rgb = SLATE_LIGHT

    paragraph.add_run("\t")

    run = paragraph.add_run("Page ")
    run.font.size = Pt(8)
    run.font.color.rgb = SLATE_LIGHT
    add_field(paragraph, "PAGE", "1")
    run2 = paragraph.add_run(" sur ")
    run2.font.size = Pt(8)
    run2.font.color.rgb = SLATE_LIGHT
    add_field(paragraph, "NUMPAGES", "1")


def add_header(section):
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("DGRAD  •  Direction des Archives et des Nouvelles Technologies")
    run.font.size = Pt(8)
    run.font.color.rgb = SLATE_LIGHT
    run.font.name = "Calibri"


def add_horizontal_rule(paragraph, color=GOLD_HEX, size="16"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def add_caption(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.font.size = Pt(8)
    run.font.italic = True
    run.font.color.rgb = SLATE_LIGHT


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.add_run(item)


def add_numbered(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Number")
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.add_run(item)


def add_callout(doc, title, text, fill=LIGHT_GOLD, border=GOLD_HEX):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color=border, size="6")
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    cell.width = Cm(16.2)

    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(title)
    run.font.bold = True
    run.font.size = Pt(9.5)
    run.font.color.rgb = NAVY

    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(2)
    run2 = p2.add_run(text)
    run2.font.size = Pt(9)
    run2.font.color.rgb = SLATE

    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_code_block(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color="E2E8F0", size="4")
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F8FAFC")
    cell.width = Cm(16.2)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(doc, headers, rows, widths=None, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_table_borders(table)

    header = table.rows[0]
    set_repeat_table_header(header)
    for index, title in enumerate(headers):
        cell = header.cells[index]
        set_cell_shading(cell, NAVY_HEX)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(1)
        run = p.add_run(title)
        run.font.bold = True
        run.font.size = Pt(font_size)
        run.font.color.rgb = WHITE
        if widths:
            cell.width = Cm(widths[index])

    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        for col_index, value in enumerate(row):
            cell = cells[col_index]
            if row_index % 2 == 1:
                set_cell_shading(cell, "F8FAFC")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            run = p.add_run(str(value))
            run.font.size = Pt(font_size)
            run.font.color.rgb = SLATE
            if widths:
                cell.width = Cm(widths[col_index])

    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_image(doc, path, width_cm, caption):
    doc.add_picture(path, width=Cm(width_cm))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_caption(doc, caption)


# ---------------------------------------------------------------------------
# Document
# ---------------------------------------------------------------------------

def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, color, before, after in [
        ("Heading 1", 18, NAVY, 20, 10),
        ("Heading 2", 13.5, NAVY, 14, 6),
        ("Heading 3", 11.5, GOLD, 10, 4),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ["List Bullet", "List Number"]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(10.5)


def build_document():
    arch_image = os.path.join(ASSETS_DIR, "architecture.png")
    data_image = os.path.join(ASSETS_DIR, "data_model.png")
    build_architecture_image(arch_image)
    build_data_model_image(data_image)

    doc = Document()
    configure_styles(doc)

    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.different_first_page_header_footer = True
    add_header(section)
    add_page_number_footer(section)

    properties = doc.core_properties
    properties.title = "Documentation technique et fonctionnelle — Plateforme de gestion des archives et des titres de perception"
    properties.subject = "Archivage ordinaire, titres de perception, OCR et gestion des accès"
    properties.author = "DGRAD — DANTIC"
    properties.keywords = "DGRAD, archives, titres de perception, OCR, Laravel, React, MySQL"
    properties.comments = "Documentation générée pour la plateforme DGRAD Archive."

    # ------------------------------------------------------------------
    # Page de garde
    # ------------------------------------------------------------------
    if os.path.exists(LOGO):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(20)
        p.add_run().add_picture(LOGO, width=Cm(3.1))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("RÉPUBLIQUE DÉMOCRATIQUE DU CONGO")
    run.font.size = Pt(9)
    run.font.bold = True
    run.font.color.rgb = SLATE_LIGHT
    p.paragraph_format.space_before = Pt(18)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("DIRECTION GÉNÉRALE DES RECETTES ADMINISTRATIVES,\nJUDICIAIRES, DOMANIALES ET DE PARTICIPATIONS")
    run.font.size = Pt(9)
    run.font.color.rgb = SLATE_LIGHT
    run.font.bold = True

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(45)
    add_horizontal_rule(p)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(28)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run("Documentation technique\net fonctionnelle")
    run.font.name = "Calibri"
    run.font.size = Pt(31)
    run.font.bold = True
    run.font.color.rgb = NAVY

    p = doc.add_paragraph()
    run = p.add_run("Plateforme de gestion des archives et des titres de perception")
    run.font.size = Pt(15)
    run.font.color.rgb = GOLD
    run.font.bold = True

    p = doc.add_paragraph()
    run = p.add_run(
        "Archivage ordinaire, titres de perception, recherche OCR, gestion multi-divisions, "
        "rôles, permissions et journalisation."
    )
    run.font.size = Pt(10.5)
    run.font.color.rgb = SLATE
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.right_indent = Cm(1.2)

    doc.add_paragraph().paragraph_format.space_after = Pt(18)

    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_borders(table, color="E2E8F0")
    table.autofit = False
    control = [
        ("Version du document", "1.0"),
        ("Date", "Octobre 2026"),
        ("Statut", "Version de référence"),
        ("Diffusion", "DGRAD — DANTIC / équipe projet"),
    ]
    for index, (label, value) in enumerate(control):
        left = table.cell(index, 0)
        right = table.cell(index, 1)
        left.width = Cm(5.2)
        right.width = Cm(9.8)
        set_cell_shading(left, "F4F6FA")
        p1 = left.paragraphs[0]
        r1 = p1.add_run(label)
        r1.font.bold = True
        r1.font.size = Pt(9)
        r1.font.color.rgb = NAVY
        p2 = right.paragraphs[0]
        r2 = p2.add_run(value)
        r2.font.size = Pt(9)
        r2.font.color.rgb = SLATE

    doc.add_paragraph().paragraph_format.space_after = Pt(30)
    p = doc.add_paragraph()
    add_horizontal_rule(p, color=GOLD_HEX, size="8")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("DGRAD  •  DANTIC  •  Gestion documentaire et financière")
    run.font.size = Pt(8)
    run.font.color.rgb = SLATE_LIGHT
    run.font.bold = True

    doc.add_page_break()

    # ------------------------------------------------------------------
    # Sommaire
    # ------------------------------------------------------------------
    p = doc.add_paragraph()
    run = p.add_run("Sommaire")
    run.font.name = "Calibri"
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = NAVY
    add_horizontal_rule(p, color=GOLD_HEX, size="12")
    p.paragraph_format.space_after = Pt(12)

    add_toc(doc)
    doc.add_page_break()

    # ------------------------------------------------------------------
    # 1. Introduction
    # ------------------------------------------------------------------
    doc.add_heading("1. Introduction", level=1)
    doc.add_heading("1.1 Contexte", level=2)
    doc.add_paragraph(
        "La Direction Générale des Recettes Administratives, Judiciaires, Domaniales et de "
        "Participations (DGRAD) gère un volume important de documents administratifs, de titres "
        "de perception et de pièces justificatives. La plateforme DGRAD Archive centralise ces "
        "activités dans une application web unique, sécurisée et traçable."
    )
    doc.add_paragraph(
        "Le projet couvre deux domaines métier complémentaires : l’archivage ordinaire, qui "
        "concerne les documents administratifs courants, et la gestion des titres de perception, "
        "qui couvre les notes, avis, amendes, extraits et états généraux émis par la DGRAD."
    )

    doc.add_heading("1.2 Objectifs", level=2)
    add_bullets(doc, [
        "Centraliser les archives et les titres de perception dans une base de données unique.",
        "Garantir la traçabilité des opérations via un journal d’audit.",
        "Gérer finement les accès par utilisateur, rôle, permission, division et direction.",
        "Permettre la recherche OCR dans les documents numérisés.",
        "Automatiser la numérotation et le contrôle des champs obligatoires des titres.",
        "Fournir des tableaux de bord adaptés aux permissions et au périmètre de chaque agent.",
    ])

    doc.add_heading("1.3 Public visé", level=2)
    add_table(doc,
              ["Public", "Utilisation du document"],
              [
                  ["Direction / décideurs", "Compréhension du périmètre, des modules et des règles de gouvernance."],
                  ["Équipe projet / DANTIC", "Maintenance, exploitation, évolution et déploiement de la plateforme."],
                  ["Administrateurs fonctionnels", "Gestion des utilisateurs, rôles, divisions et référentiels."],
                  ["Agents utilisateurs", "Guide d’utilisation des archives, titres, fichiers et recherche OCR."],
              ],
              widths=[4.2, 11.8])

    doc.add_heading("1.4 Périmètre du document", level=2)
    doc.add_paragraph(
        "Le document décrit l’architecture, le modèle de données, les modules backend et frontend, "
        "les règles de sécurité, les procédures d’installation et les principaux cas d’utilisation. "
        "Il ne remplace pas la documentation officielle des frameworks utilisés."
    )

    # ------------------------------------------------------------------
    # 2. Vue d'ensemble
    # ------------------------------------------------------------------
    doc.add_heading("2. Vue d’ensemble du projet", level=1)
    doc.add_heading("2.1 Modules fonctionnels", level=2)
    add_table(doc,
              ["Module", "Description", "Utilisateurs concernés"],
              [
                  ["Tableau de bord", "Vision consolidée, KPI, activité récente et accès rapides.", "Tous selon permissions"],
                  ["Archivage ordinaire", "Archives, types de documents, classeurs, emplacements, fichiers et OCR.", "Archivistes, agents AD"],
                  ["Titres de perception", "Titres, catégories, assujettis, centres, articles budgétaires.", "Gestionnaires titres"],
                  ["Recherche OCR", "Recherche plein texte dans les pièces numérisées.", "Agents autorisés"],
                  ["Administration", "Utilisateurs, rôles, permissions, divisions, directions, journal.", "Administrateurs"],
                  ["Profil", "Informations personnelles, avatar, mot de passe et périmètre.", "Tous les utilisateurs"],
              ],
              widths=[4.0, 9.2, 3.0])

    doc.add_heading("2.2 Acteurs et rôles", level=2)
    add_bullets(doc, [
        "Administrateur : accès complet à la plateforme, aux comptes, rôles et paramètres.",
        "Archiviste AD : gestion de l’archivage ordinaire et consultation OCR.",
        "Gestionnaire Titres : gestion des titres de perception, catégories et référentiels associés.",
        "Agent : accès aux modules autorisés par son rôle et son périmètre de divisions.",
        "Direction : accès transversal à toutes les divisions d’une direction lorsque l’option est activée.",
    ])

    doc.add_heading("2.3 Flux métier principaux", level=2)
    add_numbered(doc, [
        "Création ou import d’une archive avec classement, division, classeur et emplacement.",
        "Création d’un titre de perception selon une catégorie, avec numérotation et champs obligatoires.",
        "Ajout des pièces justificatives PDF, import ou scanner.",
        "Lancement de l’OCR et correction éventuelle du texte extrait.",
        "Validation du titre et journalisation de l’opération.",
        "Recherche et consultation selon les permissions et le périmètre de divisions.",
    ])

    # ------------------------------------------------------------------
    # 3. Architecture
    # ------------------------------------------------------------------
    doc.add_heading("3. Architecture technique", level=1)
    doc.add_heading("3.1 Architecture globale", level=2)
    add_image(doc, arch_image, 16.2, "Figure 1 — Architecture applicative de la plateforme DGRAD Archive")

    doc.add_heading("3.2 Composants", level=2)
    add_table(doc,
              ["Composant", "Responsabilité"],
              [
                  ["Frontend React", "Interface utilisateur, navigation, formulaires, tableaux de bord, contrôles d’accès visuels."],
                  ["API Laravel", "Authentification, autorisations, logique métier, validation, journalisation."],
                  ["MySQL 8", "Persistance des données, relations, index et contraintes d’intégrité."],
                  ["Service OCR Flask", "Extraction de texte depuis les PDF et images, moteur Tesseract."],
                  ["Scanner Windows", "Numérisation de documents papier et transmission des PDF à l’API."],
                  ["Stockage fichiers", "Conservation des PDF, images, scans et avatars."],
              ],
              widths=[4.0, 12.0])

    doc.add_heading("3.3 Technologies", level=2)
    add_table(doc,
              ["Couche", "Technologies"],
              [
                  ["Backend", "PHP 8, Laravel 8, Laravel Sanctum, Eloquent ORM"],
                  ["Base de données", "MySQL 8, utf8mb4, migrations Laravel"],
                  ["Frontend", "React 18, TypeScript, Vite 5, Tailwind CSS, Axios, React Router"],
                  ["OCR", "Python 3.10, Flask, Flask-RESTful, Tesseract, pdfplumber, PyPDF2"],
                  ["Outils", "Git, GitHub, Composer, npm, Windows Scanner"],
              ],
              widths=[4.0, 12.0])

    # ------------------------------------------------------------------
    # 4. Modèle de données
    # ------------------------------------------------------------------
    doc.add_heading("4. Modèle de données", level=1)
    doc.add_heading("4.1 Principes", level=2)
    doc.add_paragraph(
        "Le schéma repose sur une séparation claire entre l’organisation, les utilisateurs, les "
        "référentiels, les archives, les titres, les fichiers et l’audit. Les entités principales "
        "utilisent des clés étrangères avec des règles de suppression adaptées : restriction, "
        "mise à null ou cascade selon le niveau de criticité."
    )
    add_image(doc, data_image, 16.2, "Figure 2 — Cartographie simplifiée des domaines de données")

    doc.add_heading("4.2 Tables principales", level=2)
    add_table(doc,
              ["Domaine", "Tables"],
              [
                  ["Organisation", "directions, divisions, centre_ordonnancements"],
                  ["Utilisateurs & sécurité", "users, roles, permissions, user_role, role_permission, division_user, personal_access_tokens"],
                  ["Référentiels", "classeurs, emplacements, assujettis, types_documents, categories_titres, articles_budgetaires, mots_cles"],
                  ["Archives", "archives, archive_mot_cle"],
                  ["Titres", "titres_perception, titre_mot_cle"],
                  ["Fichiers & OCR", "fichiers_documents, traitements_ocr"],
                  ["Audit & technique", "audit_logs, migrations, jobs, failed_jobs, password_resets"],
              ],
              widths=[4.0, 12.0])

    doc.add_heading("4.3 Relations structurantes", level=2)
    add_bullets(doc, [
        "Une direction contient plusieurs divisions ; une division appartient à une seule direction.",
        "Un utilisateur possède une division principale obligatoire et peut avoir accès à plusieurs divisions via division_user.",
        "Un titre appartient obligatoirement à une catégorie et peut être rattaché à un assujetti, un centre, un article budgétaire et une division.",
        "Une archive est rattachée à une division, un type de document, un classeur et un emplacement.",
        "Les fichiers sont polymorphes : ils peuvent appartenir à une archive ou à un titre.",
        "Un fichier possède un ou plusieurs traitements OCR.",
    ])

    # ------------------------------------------------------------------
    # 5. Backend
    # ------------------------------------------------------------------
    doc.add_heading("5. Backend Laravel", level=1)
    doc.add_heading("5.1 Structure", level=2)
    add_code_block(doc,
                   "backend/\n"
                   "  app/Http/Controllers/     contrôleurs métier\n"
                   "  app/Models/               modèles Eloquent et alias de compatibilité\n"
                   "  app/Jobs/                 traitements OCR asynchrones\n"
                   "  app/Services/             service d’appel OCR\n"
                   "  app/Helpers/              journalisation applicative\n"
                   "  database/migrations/      schéma de la base\n"
                   "  database/seeders/         données initiales\n"
                   "  routes/api.php            routes de l’API\n"
                   "  config/                   configuration Laravel et OCR")

    doc.add_heading("5.2 Authentification et sécurité", level=2)
    add_bullets(doc, [
        "Authentification par Laravel Sanctum avec tokens Bearer.",
        "Mots de passe hachés avec Hash::make et vérifiés avec Hash::check.",
        "Contrôle des permissions via le middleware CheckPermission.",
        "Contrôle du périmètre de divisions dans les contrôleurs métier.",
        "Journalisation des actions dans audit_logs.",
        "Protection des routes sensibles par auth:sanctum et permission:*.",
    ])

    doc.add_heading("5.3 Modules et endpoints principaux", level=2)
    add_table(doc,
              ["Module", "Endpoints principaux"],
              [
                  ["Authentification", "POST /connexion, POST /logout"],
                  ["Utilisateurs", "GET/POST /mon-utilisateurs, PUT/DELETE /mon-utilisateurs/{id}, assign-roles, assign-directions"],
                  ["Rôles / permissions", "GET/POST /roles, /permissions, assign-permissions"],
                  ["Organisation", "GET/POST /directions, /divisions, /mes-divisions, /mes-directions"],
                  ["Archives", "GET/POST /archives, GET/PUT/DELETE /archives/{id}"],
                  ["Titres", "GET/POST /titres, GET/PUT/DELETE /titres/{id}, POST /titres/{id}/validate"],
                  ["Référentiels", "types-documents, categories-titres, mots-cles, assujettis, articles"],
                  ["Fichiers / OCR", "POST /fichiers/upload, GET /fichiers, POST /fichiers/{id}/ocr, PUT /fichiers/{id}/text"],
                  ["Recherche OCR", "GET /ocr/recherche"],
                  ["Dashboards", "GET /dashboard/statistics, /dashboard/recent, /dashboards/notes/statistics, /dashboards/notes/recent"],
                  ["Journal", "GET /journal, /journal/stats, /journal/tables"],
              ],
              widths=[4.2, 11.8])

    doc.add_heading("5.4 Règles métier des titres", level=2)
    doc.add_paragraph(
        "Chaque catégorie de titre peut définir une règle de numérotation et une liste de champs "
        "obligatoires. La règle de numérotation accepte les placeholders {YYYY}, {SEQ} et {CAT}. "
        "Les champs obligatoires utilisent les noms techniques des colonnes de titres_perception."
    )
    add_callout(
        doc,
        "Exemple",
        "Catégorie AMENDE : règle AMD-{YYYY}-{SEQ}, champs obligatoires id_assujetti. "
        "La référence générée peut être AMD-2026-000001.",
    )
    add_table(doc,
              ["Code", "Libellé", "Assujetti", "Champs obligatoires"],
              [
                  ["NOTE_PERCEPTION", "Note de perception", "Oui", "id_assujetti"],
                  ["AVIS_RECOUVREMENT", "Avis de mise en recouvrement", "Oui", "id_assujetti"],
                  ["AMENDE", "Amende", "Oui", "id_assujetti"],
                  ["EXTRAIT_ROLE", "Extrait de rôle", "Oui", "id_assujetti"],
                  ["ACQUIS_LIBERATOIRE", "Acquis libératoire", "Oui", "id_assujetti"],
                  ["RESTE_A_RECOUVRER", "Reste à recouvrer", "Non", "Aucun"],
                  ["EXTRAIT_JOURNALIER", "Extrait journalier", "Non", "Aucun"],
                  ["EXTRAIT_ACTE_SERVICE", "Extrait par acte et service", "Non", "Aucun"],
                  ["TABLEAU_RETROCESSION_SERVICE", "Tableau de rétrocession par service", "Non", "Aucun"],
              ],
              widths=[4.4, 7.6, 1.8, 2.4])

    doc.add_heading("5.5 Journalisation et audit", level=2)
    doc.add_paragraph(
        "Les actions sensibles sont enregistrées dans audit_logs avec l’utilisateur, sa division, "
        "l’action, l’entité concernée, l’identifiant, la description, les valeurs avant/après, "
        "l’adresse IP et le user-agent. Le journal est consultable depuis l’interface "
        "d’administration pour les utilisateurs disposant de la permission audit.view."
    )

    # ------------------------------------------------------------------
    # 6. Frontend
    # ------------------------------------------------------------------
    doc.add_heading("6. Application frontend", level=1)
    doc.add_heading("6.1 Structure", level=2)
    add_code_block(doc,
                   "Frontend/\n"
                   "  src/Screens/        écrans métier et pages de formulaires\n"
                   "  src/Composant/      layout, menus, composants réutilisables\n"
                   "  src/Modals/         fenêtres de confirmation, fichiers, OCR\n"
                   "  src/services/       appels API centralisés avec Axios\n"
                   "  src/types/          types TypeScript du modèle de données\n"
                   "  src/Loading/        composants de chargement et skeletons")

    doc.add_heading("6.2 Écrans principaux", level=2)
    add_table(doc,
              ["Écran", "Fonction"],
              [
                  ["Connexion", "Authentification et redirection selon les permissions."],
                  ["Dashboard", "Vue consolidée permissions + modules accessibles."],
                  ["Dashboard Archivage", "Indicateurs, top classeurs, directions et derniers versements."],
                  ["Dashboard Titres", "Indicateurs, top centres, articles et assujettis."],
                  ["Archives", "Liste, filtres, recherche OCR, fichiers et actions."],
                  ["Nouvelle archive", "Formulaire par sections, direction → division, pièces jointes."],
                  ["Titres", "Liste, filtres, validation, fichiers et actions."],
                  ["Nouveau titre", "Formulaire par catégorie, numérotation et champs obligatoires."],
                  ["Recherche OCR", "Recherche plein texte, extraits, texte complet et accès aux fiches."],
                  ["Référentiels", "Types de documents, catégories, mots-clés, assujettis, articles."],
                  ["Administration", "Utilisateurs, rôles, permissions, divisions, journal."],
                  ["Profil", "Informations, avatar, mot de passe, divisions et rôles."],
              ],
              widths=[4.2, 11.8])

    doc.add_heading("6.3 Design system", level=2)
    add_bullets(doc, [
        "Palette institutionnelle : navy, or, gris ardoise.",
        "Typographie Calibri / Manrope pour les titres et Inter pour l’interface.",
        "Cartes, tableaux, badges de statut et composants de chargement homogènes.",
        "Skeletons de chargement sur les tableaux de bord.",
        "Modals de confirmation partagés pour toutes les suppressions.",
        "Formulaires en pages dédiées pour les archives, titres et utilisateurs.",
    ])

    doc.add_heading("6.4 Permissions côté interface", level=2)
    doc.add_paragraph(
        "Les menus, boutons, modules et tableaux de bord sont conditionnés par les permissions "
        "retournées à la connexion. Les appels API restent toutefois protégés côté backend : "
        "masquer un élément dans l’interface ne remplace jamais le contrôle serveur."
    )

    # ------------------------------------------------------------------
    # 7. Utilisateurs, rôles et divisions
    # ------------------------------------------------------------------
    doc.add_heading("7. Gestion des utilisateurs, rôles et divisions", level=1)
    doc.add_heading("7.1 Comptes utilisateurs", level=2)
    add_bullets(doc, [
        "Chaque utilisateur possède un nom, un prénom, un email unique, un mot de passe et une division principale.",
        "Le statut peut être active, inactive ou bloque.",
        "Un utilisateur peut recevoir plusieurs rôles.",
        "Le profil permet de modifier les informations, le téléphone, l’avatar et le mot de passe.",
    ])

    doc.add_heading("7.2 Rôles et permissions", level=2)
    add_table(doc,
              ["Rôle", "Portée principale"],
              [
                  ["Admin", "Tous les modules et toutes les divisions."],
                  ["Archiviste AD", "Archivage ordinaire, référentiels AD, fichiers, OCR, rapports."],
                  ["Gestionnaire Titres", "Titres de perception, catégories, centres, assujettis, articles, OCR."],
              ],
              widths=[4.5, 11.5])
    doc.add_paragraph(
        "Les permissions sont regroupées par module : dashboard, archives, titres, documents, "
        "administration, rapports. Les codes sont contrôlés par le middleware CheckPermission."
    )

    doc.add_heading("7.3 Multi-divisions et accès direction", level=2)
    doc.add_paragraph(
        "La division principale reste obligatoire. Des divisions supplémentaires peuvent être "
        "accordées via la table division_user. L’option « Accès à toute la direction » étend "
        "automatiquement l’accès à toutes les divisions de la direction de rattachement."
    )
    add_callout(
        doc,
        "Cas d’usage — Directeur",
        "Un directeur rattaché à la division INFO de la direction DANTIC peut consulter les "
        "archives et les titres de toutes les divisions DANTIC en activant l’accès direction.",
    )

    doc.add_heading("7.4 Portée des données", level=2)
    doc.add_paragraph(
        "Les listes, tableaux de bord et recherches appliquent le périmètre de divisions de "
        "l’utilisateur. L’administrateur conserve une vision globale. Les endpoints sensibles "
        "vérifient également les permissions : archives.view, titres.view, documents.view, "
        "ocr.view, users.view, audit.view."
    )

    # ------------------------------------------------------------------
    # 8. OCR et fichiers
    # ------------------------------------------------------------------
    doc.add_heading("8. OCR, fichiers et scanner", level=1)
    doc.add_heading("8.1 Cycle de vie d’une pièce", level=2)
    add_numbered(doc, [
        "Import d’un PDF ou d’une image autorisée (PDF, JPG, PNG, TIFF, WEBP).",
        "Enregistrement dans fichiers_documents avec métadonnées et hash SHA-256.",
        "Création d’un traitement OCR en attente.",
        "Extraction du texte par le service Flask / Tesseract.",
        "Stockage du texte dans traitements_ocr.texte_extrait.",
        "Correction manuelle éventuelle et recherche plein texte.",
    ])
    doc.add_paragraph(
        "Les fichiers sont polymorphes : entite_type vaut archive ou titre_perception. "
        "Ils sont rattachés à l’entité via entite_id."
    )

    doc.add_heading("8.2 Service OCR", level=2)
    add_bullets(doc, [
        "Service Python Flask exposant /api/ocr, /api/ocr-batch et /api/health.",
        "Moteur Tesseract pour les PDF et images.",
        "Appel sécurisé par jeton Bearer.",
        "Configuration backend : FLASK_OCR_URL et FLASK_OCR_BATCH_URL.",
        "Le bouton « Lancer OCR » exécute le traitement et met à jour le texte.",
    ])
    add_callout(
        doc,
        "Dépannage",
        "Si le service distant renvoie une page Coming Soon ou 404, démarrer le service local "
        "avec ocr/start_ocr.bat puis configurer FLASK_OCR_URL=http://127.0.0.1:5000/api/ocr.",
    )

    doc.add_heading("8.3 Recherche OCR", level=2)
    doc.add_paragraph(
        "La recherche OCR interroge traitements_ocr.texte_extrait et retourne les entités "
        "correspondantes avec un extrait contextualisé. Les résultats sont limités au périmètre "
        "de divisions de l’utilisateur et nécessitent la permission ocr.view."
    )

    doc.add_heading("8.4 Scanner", level=2)
    doc.add_paragraph(
        "Le scanner Windows transmet les PDF à l’API. Le backend rattache automatiquement le "
        "scan à l’archive ou au titre courant, puis déclenche l’OCR. L’interface propose "
        "l’import et le scan dans le modal Fichiers & OCR."
    )

    # ------------------------------------------------------------------
    # 9. Tableaux de bord
    # ------------------------------------------------------------------
    doc.add_heading("9. Tableaux de bord", level=1)
    doc.add_heading("9.1 Dashboard unifié", level=2)
    doc.add_paragraph(
        "Vue consolidée des modules autorisés, des KPI, de l’activité récente et des raccourcis. "
        "Les requêtes sont conditionnées par les permissions pour éviter tout appel non autorisé."
    )
    doc.add_heading("9.2 Dashboard Archivage", level=2)
    doc.add_paragraph(
        "Indicateurs sur le fonds documentaire, les documents actifs, les archives et l’activité "
        "du jour. Répartition par classeur et par direction, avec derniers versements."
    )
    doc.add_heading("9.3 Dashboard Titres", level=2)
    doc.add_paragraph(
        "Indicateurs sur les titres émis, les titres actifs, les brouillons, les centres, "
        "articles et assujettis. Le montant n’est plus mis en avant dans les indicateurs "
        "principaux."
    )
    doc.add_heading("9.4 Filtrage par permissions et divisions", level=2)
    doc.add_paragraph(
        "Chaque tableau de bord contrôle les permissions requises et applique le périmètre "
        "de divisions. Un écran d’accès refusé est affiché lorsqu’un droit manque."
    )

    # ------------------------------------------------------------------
    # 10. Installation
    # ------------------------------------------------------------------
    doc.add_heading("10. Installation et configuration", level=1)
    doc.add_heading("10.1 Prérequis", level=2)
    add_table(doc,
              ["Outil", "Version recommandée"],
              [
                  ["PHP", "8.0 ou supérieur"],
                  ["Composer", "2.x"],
                  ["MySQL", "8.0"],
                  ["Node.js", "18 ou supérieur"],
                  ["Python", "3.10"],
                  ["Tesseract OCR", "Installation Windows"],
              ],
              widths=[5.0, 11.0])

    doc.add_heading("10.2 Backend", level=2)
    add_code_block(doc,
                   "cd backend\n"
                   "composer install\n"
                   "cp .env.example .env\n"
                   "php artisan key:generate\n"
                   "php artisan migrate --seed\n"
                   "php artisan serve")

    doc.add_heading("10.3 Frontend", level=2)
    add_code_block(doc,
                   "cd Frontend\n"
                   "npm install\n"
                   "npm run dev")
    doc.add_paragraph(
        "Configurer VITE_API_URL dans le fichier d’environnement du frontend."
    )

    doc.add_heading("10.4 OCR", level=2)
    add_code_block(doc,
                   "cd ocr\n"
                   "pip install -r requirements.txt\n"
                   "python app.py")
    doc.add_paragraph("Sous Windows, le script ocr/start_ocr.bat démarre le service local.")

    doc.add_heading("10.5 Variables d’environnement principales", level=2)
    add_code_block(doc,
                   "APP_URL=http://localhost:8000\n"
                   "DB_CONNECTION=mysql\n"
                   "DB_DATABASE=bdarchive_titres_perception\n"
                   "QUEUE_CONNECTION=database\n"
                   "FLASK_OCR_URL=http://127.0.0.1:5000/api/ocr\n"
                   "FLASK_OCR_BATCH_URL=http://127.0.0.1:5000/api/ocr-batch\n"
                   "FLASK_OCR_TOKEN=ocr_secret_token_2026")

    # ------------------------------------------------------------------
    # 11. Guide d'utilisation
    # ------------------------------------------------------------------
    doc.add_heading("11. Guide d’utilisation", level=1)
    doc.add_heading("11.1 Connexion", level=2)
    doc.add_paragraph(
        "Accéder à la page de connexion, saisir l’email professionnel et le mot de passe. "
        "La redirection dépend des permissions : dashboard, archives, titres ou profil."
    )
    doc.add_heading("11.2 Créer une archive", level=2)
    add_numbered(doc, [
        "Ouvrir Archives puis « Nouvelle archive ».",
        "Saisir le titre, l’objet et le type de document.",
        "Choisir la direction puis la division.",
        "Renseigner le classeur, l’emplacement et les dates.",
        "Enregistrer, puis importer le PDF ou scanner le document.",
    ])
    doc.add_heading("11.3 Créer un titre de perception", level=2)
    add_numbered(doc, [
        "Ouvrir Titres puis « Nouveau titre ».",
        "Sélectionner la catégorie ; les champs obligatoires sont indiqués.",
        "Renseigner l’assujetti, le montant et les références selon la catégorie.",
        "Choisir la direction puis la division émettrice.",
        "Enregistrer, joindre les pièces, lancer l’OCR puis valider le titre.",
    ])
    doc.add_heading("11.4 Recherche OCR", level=2)
    doc.add_paragraph(
        "Saisir un mot-clé, une date ou une référence. Filtrer éventuellement sur les archives "
        "ou les titres. Les extraits mettent en évidence le terme recherché ; le texte complet "
        "est consultable et copiable."
    )
    doc.add_heading("11.5 Administration", level=2)
    add_bullets(doc, [
        "Créer un utilisateur en choisissant direction, division, rôles et accès.",
        "Accorder des divisions supplémentaires ou l’accès à toute la direction.",
        "Gérer les rôles et permissions.",
        "Consulter le journal d’audit.",
    ])

    # ------------------------------------------------------------------
    # 12. Exploitation
    # ------------------------------------------------------------------
    doc.add_heading("12. Exploitation et maintenance", level=1)
    doc.add_heading("12.1 Commandes utiles", level=2)
    add_code_block(doc,
                   "php artisan migrate\n"
                   "php artisan db:seed --class=ReferentielSeeder\n"
                   "php artisan config:clear\n"
                   "php artisan route:clear\n"
                   "php artisan queue:work\n"
                   "npm run build")

    doc.add_heading("12.2 Sauvegarde", level=2)
    add_bullets(doc, [
        "Sauvegarder la base MySQL bdarchive_titres_perception.",
        "Sauvegarder le dossier backend/storage/app (avatars, scans, PDF).",
        "Conserver les fichiers .env de production de manière sécurisée.",
        "Tester périodiquement la restauration.",
    ])

    doc.add_heading("12.3 Dépannage", level=2)
    add_table(doc,
              ["Symptôme", "Piste de résolution"],
              [
                  ["Erreur OCR 404", "Vérifier FLASK_OCR_URL, démarrer le service OCR, vider le cache de configuration."],
                  ["OCR bloqué en attente", "Lancer le traitement manuellement ou démarrer un worker de queue."],
                  ["Utilisateur sans accès", "Vérifier le rôle, la permission et les divisions accessibles."],
                  ["Avatar non affiché", "Vérifier le fichier dans storage/app/avatar et la route publique /api/profil/avatar."],
                  ["Dashboard vide", "Vérifier les permissions et le périmètre de divisions."],
              ],
              widths=[5.0, 11.0])

    doc.add_heading("12.4 Bonnes pratiques de sécurité", level=2)
    add_bullets(doc, [
        "Ne jamais versionner les fichiers .env.",
        "Utiliser des mots de passe forts et renouveler les jetons.",
        "Limiter les privilèges au strict nécessaire.",
        "Activer HTTPS en production.",
        "Contrôler régulièrement le journal d’audit.",
    ])

    # ------------------------------------------------------------------
    # 13. Annexes
    # ------------------------------------------------------------------
    doc.add_heading("13. Annexes", level=1)
    doc.add_heading("A. Statuts", level=2)
    add_table(doc,
              ["Domaine", "Statuts"],
              [
                  ["Archives", "brouillon, actif, archive, inactif"],
                  ["Titres", "brouillon, actif, archive, annule"],
                  ["OCR", "en_attente, en_cours, termine, erreur"],
                  ["Utilisateurs", "active, inactive, bloque"],
              ],
              widths=[4.0, 12.0])

    doc.add_heading("B. Permissions principales", level=2)
    add_table(doc,
              ["Module", "Exemples de permissions"],
              [
                  ["Archives", "archives.view, archives.create, archives.update, archives.delete"],
                  ["Titres", "titres.view, titres.create, titres.update, titres.delete, titres.validate"],
                  ["Documents / OCR", "documents.view, documents.upload, documents.download, documents.delete, ocr.view, ocr.run"],
                  ["Administration", "users.view, users.create, users.update, users.delete, roles.*, permissions.*, divisions.*"],
                  ["Audit", "audit.view"],
              ],
              widths=[4.2, 11.8])

    doc.add_heading("C. Glossaire", level=2)
    add_table(doc,
              ["Terme", "Définition"],
              [
                  ["DGRAD", "Direction Générale des Recettes Administratives, Judiciaires, Domaniales et de Participations."],
                  ["DANTIC", "Direction des Archives et des Nouvelles Technologies."],
                  ["Archive", "Document administratif conservé et classé."],
                  ["Titre de perception", "Document constatant une créance ou un droit à recouvrer."],
                  ["Assujetti", "Personne physique ou morale redevable."],
                  ["OCR", "Reconnaissance optique de caractères."],
                  ["Périmètre", "Ensemble des divisions accessibles à un utilisateur."],
                  ["RBAC", "Contrôle d’accès basé sur les rôles."],
              ],
              widths=[4.2, 11.8])

    doc.add_heading("D. Références", level=2)
    add_bullets(doc, [
        "Laravel 8 — https://laravel.com/docs/8.x",
        "React — https://react.dev",
        "Vite — https://vitejs.dev",
        "MySQL 8 — https://dev.mysql.com/doc",
        "Tesseract OCR — https://github.com/tesseract-ocr/tesseract",
    ])

    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    path = build_document()
    print(path)
